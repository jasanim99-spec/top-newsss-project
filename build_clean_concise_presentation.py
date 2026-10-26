import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Academic-Corporate Modern Color Palette
    NAVY = RGBColor(15, 30, 65)          # #0F1E41
    BLUE = RGBColor(28, 78, 155)         # #1C4E9B
    LIGHT_BLUE = RGBColor(238, 244, 255) # Soft Blue Tint
    LIGHT_BG = RGBColor(248, 249, 252)   # #F8F9FC Canvas
    CARD_BG = RGBColor(255, 255, 255)    # White
    CARD_BORDER = RGBColor(222, 228, 238)# Light Gray Border
    TEXT_DARK = RGBColor(17, 24, 39)     # #111827 Charcoal Dark
    TEXT_MUTED = RGBColor(100, 116, 139) # #64748B Slate Muted
    GREEN = RGBColor(16, 149, 93)        # #10955D Success Green
    GREEN_LIGHT = RGBColor(236, 253, 245)
    GOLD = RGBColor(217, 119, 6)         # #D97706 Amber Gold
    GOLD_LIGHT = RGBColor(254, 243, 199)
    RED = RGBColor(220, 38, 38)
    RED_LIGHT = RGBColor(254, 226, 226)

    # Asset paths
    PPSU_LOGO = "poster_assets/page_0_Image13.png"
    SOE_LOGO = "poster_assets/page_0_Image15.png"
    IMG_HERO = "ppt_extracted_assets/page_52_img_0.jpeg"
    IMG_FLOW = "ppt_extracted_assets/page_40_img_0.jpeg"
    IMG_ARCH = "ppt_extracted_assets/page_39_img_0.jpeg"
    IMG_VIDEO = "ppt_extracted_assets/page_72_img_0.jpeg"
    IMG_ADMIN = "ppt_extracted_assets/page_53_img_0.jpeg"
    IMG_PRESS = "ppt_extracted_assets/page_73_img_0.png"
    IMG_QR = "ppt_extracted_assets/github_repo_qr.png"

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.fill.background()
        return bg

    def add_top_banner(slide):
        header_rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.85))
        header_rect.fill.solid()
        header_rect.fill.fore_color.rgb = CARD_BG
        header_rect.line.fill.background()

        if os.path.exists(PPSU_LOGO):
            slide.shapes.add_picture(PPSU_LOGO, Inches(0.6), Inches(0.1), height=Inches(0.65))

        if os.path.exists(SOE_LOGO):
            slide.shapes.add_picture(SOE_LOGO, Inches(10.7), Inches(0.12), height=Inches(0.62))

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0.84), Inches(13.333), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

    def add_slide_header_block(slide, slide_num, title_text, subtitle_text):
        cat_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.94), Inches(12.133), Inches(0.24))
        tf = cat_box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.text = f"P P SAVANI UNIVERSITY  ·  SCHOOL OF ENGINEERING  ·  FINAL INTERNSHIP PRESENTATION (2026–27)  |  SLIDE {slide_num}"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = BLUE

        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.18), Inches(12.133), Inches(0.40))
        tf2 = title_box.text_frame
        tf2.word_wrap = True
        tf2.margin_top = tf2.margin_bottom = tf2.margin_left = tf2.margin_right = 0
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = NAVY

        sub_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.58), Inches(12.133), Inches(0.28))
        tf3 = sub_box.text_frame
        tf3.word_wrap = True
        tf3.margin_top = tf3.margin_bottom = tf3.margin_left = tf3.margin_right = 0
        p3 = tf3.paragraphs[0]
        p3.text = subtitle_text
        p3.font.size = Pt(12)
        p3.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: TITLE SLIDE (Cover)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    add_top_banner(s1)

    left_c = add_card(s1, Inches(0.6), Inches(1.15), Inches(7.5), Inches(5.95))

    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.35), Inches(4.8), Inches(0.35))
    badge.fill.solid()
    badge.fill.fore_color.rgb = LIGHT_BLUE
    badge.line.color.rgb = BLUE
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "FINAL YEAR INTERNSHIP PROJECT PRESENTATION · 2026–27"
    p_b.font.size = Pt(10.5)
    p_b.font.bold = True
    p_b.font.color.rgb = BLUE
    p_b.alignment = PP_ALIGN.CENTER

    t_box = s1.shapes.add_textbox(Inches(0.9), Inches(1.80), Inches(6.9), Inches(1.8))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p1 = tf_t.paragraphs[0]
    p1.text = "TOP NEWS PORTAL"
    p1.font.size = Pt(34)
    p1.font.bold = True
    p1.font.color.rgb = NAVY

    p2 = tf_t.add_paragraph()
    p2.text = "Full-Stack Digital Media & Short-Video Streaming Platform with PostgreSQL"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = BLUE
    p2.space_before = Pt(4)

    p3 = tf_t.add_paragraph()
    p3.text = "Decoupled 3-tier architecture unifying news reading, vertical reels & newsroom governance."
    p3.font.size = Pt(11.5)
    p3.font.color.rgb = TEXT_DARK
    p3.space_before = Pt(4)

    badges_box = s1.shapes.add_textbox(Inches(0.9), Inches(3.68), Inches(6.9), Inches(0.35))
    tf_bd = badges_box.text_frame
    p_bd = tf_bd.paragraphs[0]
    p_bd.text = "✓ Live on Vercel CDN    ✓ Neon PostgreSQL 16    ✓ React 18 & TypeScript 5.8"
    p_bd.font.size = Pt(11.5)
    p_bd.font.bold = True
    p_bd.font.color.rgb = GREEN

    cand_box = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(4.15), Inches(6.9), Inches(2.65))
    cand_box.fill.solid()
    cand_box.fill.fore_color.rgb = LIGHT_BG
    cand_box.line.color.rgb = CARD_BORDER
    cand_box.line.width = Pt(1)
    tf_c = cand_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = Inches(0.25)
    tf_c.margin_top = Inches(0.15)

    pc1 = tf_c.paragraphs[0]
    pc1.text = "STUDENT & PROJECT PARTICULARS"
    pc1.font.size = Pt(11)
    pc1.font.bold = True
    pc1.font.color.rgb = BLUE

    pc2 = tf_c.add_paragraph()
    pc2.text = "Student: JASANI MEET PARESHBHAI   |   Enrollment: 23SE02ML077"
    pc2.font.size = Pt(13)
    pc2.font.bold = True
    pc2.font.color.rgb = NAVY
    pc2.space_before = Pt(4)

    pc3 = tf_c.add_paragraph()
    pc3.text = "Degree: B.Tech Computer Science Engineering (ML & AI) — Semester 7"
    pc3.font.size = Pt(11.5)
    pc3.font.color.rgb = TEXT_DARK
    pc3.space_before = Pt(2)

    pc4 = tf_c.add_paragraph()
    pc4.text = "Company Mentor: Ms. Priya Trivedi, Octfis Techno LLP, Surat (18 Weeks)"
    pc4.font.size = Pt(11)
    pc4.font.color.rgb = TEXT_MUTED
    pc4.space_before = Pt(2)

    pc5 = tf_c.add_paragraph()
    pc5.text = "Institute Mentor: Mr. Rakesh Tanwar, Assistant Professor, School of Engineering"
    pc5.font.size = Pt(11)
    pc5.font.color.rgb = TEXT_MUTED

    pc6 = tf_c.add_paragraph()
    pc6.text = "School of Engineering, P P Savani University, Surat, Gujarat"
    pc6.font.size = Pt(11.5)
    pc6.font.bold = True
    pc6.font.color.rgb = NAVY
    pc6.space_before = Pt(3)

    right_c = add_card(s1, Inches(8.35), Inches(1.15), Inches(4.4), Inches(5.95))
    if os.path.exists(IMG_HERO):
        s1.shapes.add_picture(IMG_HERO, Inches(8.55), Inches(1.35), width=Inches(4.0))

    cap_box = s1.shapes.add_textbox(Inches(8.55), Inches(3.60), Inches(4.0), Inches(3.3))
    tf_cap = cap_box.text_frame
    tf_cap.word_wrap = True
    p_cap = tf_cap.paragraphs[0]
    p_cap.text = "CORE DELIVERABLES AT A GLANCE"
    p_cap.font.size = Pt(13)
    p_cap.font.bold = True
    p_cap.font.color.rgb = NAVY

    delivs = [
        "• Decoupled 3-Tier Web Stack (Vite + Express + Neon DB)",
        "• Multi-Lingual Feeds (Gujarati & English) + Voice Search",
        "• Mobile-First Vertical Video Reels Player (<140ms buffer)",
        "• Newsroom Governance: Editorial Queue & QR Press Passes",
        "• Native VAPID Web Push Alerts (<1s delivery to browsers)",
        "• 100% Type-Safe TypeScript 5.8 & Sub-65ms API Response"
    ]
    for d in delivs:
        pd = tf_cap.add_paragraph()
        pd.text = d
        pd.font.size = Pt(11)
        pd.font.color.rgb = TEXT_DARK
        pd.space_before = Pt(3)

    # =========================================================================
    # SLIDE 2: INTRODUCTION & OVERVIEW
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_top_banner(s2)
    add_slide_header_block(s2, 2, "INTRODUCTION & OVERVIEW", "Modern Digital Journalism Meets High-Performance Cloud Architecture")

    intro_cards = [
        ("The Enterprise Challenge", [
            "• Database Lockups: Legacy PHP/WordPress sites suffer severe locks under breaking news.",
            "• High Latency: Monolithic server rendering causes slow response times (>350ms).",
            "• Fragmented Content: Text news and short-video reels are trapped in separate tools."
        ], RED, Inches(0.7), Inches(2.05)),

        ("The Decoupled 3-Tier Solution", [
            "• Client SPAs: Dual React 18 & Vite apps with TypeScript 5.8 static typing.",
            "• REST API Server: Node.js & Express engine with pg.Pool connection pooling.",
            "• Cloud Database: ACID-compliant PostgreSQL 16 hosted on Neon Cloud."
        ], BLUE, Inches(6.8), Inches(2.05)),

        ("Real-Time Media & Push Engine", [
            "• Vertical Video Reels: Swipeable mobile player for HTML5 MP4 & YouTube Shorts.",
            "• Socket.IO Sync: Instant bi-directional live news updates across all reader tabs.",
            "• Native Web Push: Zero-cost browser alerts using cryptographic VAPID keypairs."
        ], GOLD, Inches(0.7), Inches(4.65)),

        ("Internship & Production Hosting", [
            "• 18-Week Full-Stack Internship: Carried out at Octfis Techno LLP, Surat.",
            "• 100% Test Pass Rate: 36/36 automated & manual scenarios thoroughly verified.",
            "• Live Cloud Deployment: Hosted on Vercel Edge CDN with verified <65ms latency."
        ], GREEN, Inches(6.8), Inches(4.65))
    ]

    for title, points, color, left_pos, top_pos in intro_cards:
        c_box = add_card(s2, left_pos, top_pos, Inches(5.8), Inches(2.45))
        
        t_box = s2.shapes.add_textbox(left_pos + Inches(0.25), top_pos + Inches(0.18), Inches(5.3), Inches(2.1))
        tf = t_box.text_frame
        tf.word_wrap = True
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(15)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        for pt in points:
            p_b = tf.add_paragraph()
            p_b.text = pt
            p_b.font.size = Pt(12)
            p_b.font.color.rgb = TEXT_DARK
            p_b.space_before = Pt(5)

    # =========================================================================
    # SLIDE 3: PROBLEM STATEMENT & INDUSTRY BENCHMARKS
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_top_banner(s3)
    add_slide_header_block(s3, 3, "PROBLEM STATEMENT & INDUSTRY BENCHMARKS", "Why Traditional Digital News Platforms & CMS Fall Short")

    c_left = add_card(s3, Inches(0.6), Inches(1.95), Inches(7.4), Inches(5.2))

    probs = [
        ("1. Monolithic SSR Overhead & Latency Spikes",
         "Legacy PHP/WordPress sites rebuild entire HTML pages on every request. Breaking news spikes trigger unindexed database locks and slow latencies (>350ms).",
         RED),
        ("2. Fragmented Content & Short-Video Workflows",
         "Modern audiences demand vertical swipeable video reels. Traditional news sites isolate videos in external iframes without caching or view analytics.",
         GOLD),
        ("3. Inflexible Schemas & Fragmented Newsroom Desks",
         "NoSQL and document stores lack strict relational ACID integrity needed for multi-tier draft reviews, QR press passes, and reporter leave approval grids.",
         NAVY)
    ]

    for idx, (p_title, p_desc, p_col) in enumerate(probs):
        top_y = Inches(2.15) + idx * Inches(1.58)
        p_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), top_y, Inches(6.9), Inches(1.42))
        p_box.fill.solid()
        p_box.fill.fore_color.rgb = LIGHT_BG
        p_box.line.color.rgb = CARD_BORDER
        p_box.line.width = Pt(1)

        tb = s3.shapes.add_textbox(Inches(1.05), top_y + Inches(0.12), Inches(6.5), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = p_title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = p_col

        p2 = tf.add_paragraph()
        p2.text = p_desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(4)

    c_right = add_card(s3, Inches(8.25), Inches(1.95), Inches(4.5), Inches(5.2))

    tb_r = s3.shapes.add_textbox(Inches(8.45), Inches(2.10), Inches(4.1), Inches(0.35))
    tf_r = tb_r.text_frame
    p_r = tf_r.paragraphs[0]
    p_r.text = "OPERATIONAL COMPARISON"
    p_r.font.size = Pt(14)
    p_r.font.bold = True
    p_r.font.color.rgb = BLUE
    p_r.alignment = PP_ALIGN.CENTER

    # Legacy Box
    box_leg = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.45), Inches(2.55), Inches(4.1), Inches(2.05))
    box_leg.fill.solid()
    box_leg.fill.fore_color.rgb = RED_LIGHT
    box_leg.line.color.rgb = RED
    box_leg.line.width = Pt(1)
    tf_l = box_leg.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.15)
    tf_l.margin_top = Inches(0.12)
    pl1 = tf_l.paragraphs[0]
    pl1.text = "Legacy Monolithic CMS (PHP/WordPress)"
    pl1.font.size = Pt(12)
    pl1.font.bold = True
    pl1.font.color.rgb = RED

    l_bullets = [
        "• API Latency: 350 – 500 ms (high server load)",
        "• UI Stability: Full page reloads & screen flickering",
        "• Video Reels: Third-party iframes, no tracking",
        "• Push Alerts: Slow third-party plugins (4–8s)"
    ]
    for b in l_bullets:
        pb = tf_l.add_paragraph()
        pb.text = b
        pb.font.size = Pt(11)
        pb.font.color.rgb = TEXT_DARK
        pb.space_before = Pt(2)

    # TOP NEWS Box
    box_tn = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.45), Inches(4.80), Inches(4.1), Inches(2.15))
    box_tn.fill.solid()
    box_tn.fill.fore_color.rgb = GREEN_LIGHT
    box_tn.line.color.rgb = GREEN
    box_tn.line.width = Pt(1)
    tf_tn = box_tn.text_frame
    tf_tn.word_wrap = True
    tf_tn.margin_left = Inches(0.15)
    tf_tn.margin_top = Inches(0.12)
    pt1 = tf_tn.paragraphs[0]
    pt1.text = "TOP NEWS Cloud Platform"
    pt1.font.size = Pt(12)
    pt1.font.bold = True
    pt1.font.color.rgb = GREEN

    t_bullets = [
        "• API Latency: 65 ms average (28 ms indexed)",
        "• UI Stability: 100% Stable (React Query caching)",
        "• Video Reels: Integrated player with atomic SQL views",
        "• Push Alerts: 0.92 s native VAPID Web Push alerts",
        "• Reliability: Zero-flicker UI & 0 TypeScript errors"
    ]
    for b in t_bullets:
        pb = tf_tn.add_paragraph()
        pb.text = b
        pb.font.size = Pt(11)
        pb.font.color.rgb = TEXT_DARK
        pb.space_before = Pt(2)

    # =========================================================================
    # SLIDE 4: PROJECT OBJECTIVES
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_top_banner(s4)
    add_slide_header_block(s4, 4, "PROJECT OBJECTIVES", "Six Core Measurable Technical Goals Delivered & Validated")

    c_main = add_card(s4, Inches(0.6), Inches(1.95), Inches(12.133), Inches(5.2))

    objs = [
        ("#1", "Decoupled 3-Tier Architecture", [
            "• Independent React 18 / Vite SPAs & Express REST API.",
            "• Serverless Neon PostgreSQL 16 with connection pooling."
        ], BLUE),
        ("#2", "1-Click Hero Story Engine", [
            "• Instant article promotion to homepage carousel.",
            "• Automated single-active reset & Socket.IO real-time sync."
        ], GREEN),
        ("#3", "CORS-Safe Short-Video Reels", [
            "• Vertical reels player for HTML5 MP4 & YouTube Shorts.",
            "• Atomic SQL view counter increments on every video view."
        ], GOLD),
        ("#4", "Newsroom Editorial Governance", [
            "• Role-based desks: Super Admin, Chief Editor & Reporter.",
            "• Formal review queue with revision notes & leave grid."
        ], NAVY),
        ("#5", "Digital Press Card Generator", [
            "• Dynamic HTML5 Canvas produces 2x retina press passes.",
            "• Live QR code verifies reporter credentials in public."
        ], BLUE),
        ("#6", "Web Push & Ad Monetization", [
            "• VAPID crypto-signed breaking news alerts in <1 second.",
            "• Banner ad manager with active/pending status & click logs."
        ], GREEN)
    ]

    for idx, (num, title, points, color) in enumerate(objs):
        row = idx // 3
        col = idx % 3
        c_x = Inches(0.85) + col * Inches(3.95)
        c_y = Inches(2.15) + row * Inches(2.45)
        card_w = Inches(3.8)
        card_h = Inches(2.28)

        card_o = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, c_y, card_w, card_h)
        card_o.fill.solid()
        card_o.fill.fore_color.rgb = LIGHT_BG
        card_o.line.color.rgb = CARD_BORDER
        card_o.line.width = Pt(1)

        pill = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(0.15), c_y + Inches(0.15), Inches(0.7), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = color
        pill.line.fill.background()
        tf_p = pill.text_frame
        p_p = tf_p.paragraphs[0]
        p_p.text = num
        p_p.font.size = Pt(11)
        p_p.font.bold = True
        p_p.font.color.rgb = CARD_BG
        p_p.alignment = PP_ALIGN.CENTER

        t_head = s4.shapes.add_textbox(c_x + Inches(0.95), c_y + Inches(0.12), card_w - Inches(1.05), Inches(0.38))
        tf_h = t_head.text_frame
        tf_h.word_wrap = True
        ph = tf_h.paragraphs[0]
        ph.text = title
        ph.font.size = Pt(13)
        ph.font.bold = True
        ph.font.color.rgb = NAVY

        t_body = s4.shapes.add_textbox(c_x + Inches(0.15), c_y + Inches(0.58), card_w - Inches(0.3), card_h - Inches(0.68))
        tf_b = t_body.text_frame
        tf_b.word_wrap = True
        for pt in points:
            pb = tf_b.add_paragraph()
            pb.text = pt
            pb.font.size = Pt(11.5)
            pb.font.color.rgb = TEXT_DARK
            pb.space_before = Pt(3)

    # =========================================================================
    # SLIDE 5: TECHNOLOGY & ARCHITECTURE
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_top_banner(s5)
    add_slide_header_block(s5, 5, "TECHNOLOGY & ARCHITECTURE", "Decoupled 3-Tier Enterprise Architecture: Scalable, Fast & Secure")

    c_main = add_card(s5, Inches(0.6), Inches(1.95), Inches(12.133), Inches(5.2))

    tech_tiers = [
        ("Presentation Layer (Client SPAs)", [
            "• React 18.3 & Vite 5.4: Sub-second HMR & optimized client bundle.",
            "• TypeScript 5.8: 100% static type safety (0 build errors verified).",
            "• Tailwind CSS 3.4: Utility styling, glassmorphism cards & dark mode.",
            "• TanStack React Query v5: Stale-while-revalidate caching (no flicker).",
            "• HTML5 Canvas & Web Speech: Native voice search & QR press cards."
        ], BLUE, Inches(0.85)),

        ("Application Layer (Express Engine)", [
            "• Node.js 20 & Express 4.19: High-concurrency RESTful API engine.",
            "• Socket.IO v4.8: Live WebSocket broadcasting across all tabs.",
            "• VAPID Web Push: Instant push notification dispatches (<1s delivery).",
            "• Firebase JWT Auth: Secure Bearer tokens & role route guards.",
            "• Process Launcher: Custom start.js orchestrates all 3 services."
        ], NAVY, Inches(4.80)),

        ("Persistence & Cloud Storage", [
            "• Neon PostgreSQL 16: Serverless database with SSL connection pooling.",
            "• pg.Pool: Handles 500+ concurrent connections without exhaustion.",
            "• B-Tree Indexing: Achieves sub-28ms queries on slug, category & status.",
            "• 9 Normalized Tables: Strict ACID compliance & zero data drift.",
            "• Firebase Storage CDN: High-speed edge CDN for images & MP4 reels."
        ], GREEN, Inches(8.75))
    ]

    for title, points, color, left_pos in tech_tiers:
        t_card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(2.15), Inches(3.8), Inches(4.75))
        t_card.fill.solid()
        t_card.fill.fore_color.rgb = LIGHT_BG
        t_card.line.color.rgb = CARD_BORDER
        t_card.line.width = Pt(1)

        h_box = s5.shapes.add_textbox(left_pos + Inches(0.18), Inches(2.28), Inches(3.44), Inches(0.45))
        tf_h = h_box.text_frame
        tf_h.word_wrap = True
        ph = tf_h.paragraphs[0]
        ph.text = title
        ph.font.size = Pt(14)
        ph.font.bold = True
        ph.font.color.rgb = color

        b_box = s5.shapes.add_textbox(left_pos + Inches(0.18), Inches(2.80), Inches(3.44), Inches(3.95))
        tf_b = b_box.text_frame
        tf_b.word_wrap = True
        for pt in points:
            pb = tf_b.add_paragraph()
            pb.text = pt
            pb.font.size = Pt(12)
            pb.font.color.rgb = TEXT_DARK
            pb.space_before = Pt(8)

    # =========================================================================
    # SLIDE 6: PROCESS WORKFLOW
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_top_banner(s6)
    add_slide_header_block(s6, 6, "PROCESS WORKFLOW", "From News Drafting & Review to Instant Live Broadcasting")

    c_left = add_card(s6, Inches(0.6), Inches(1.95), Inches(5.8), Inches(5.2))
    if os.path.exists(IMG_FLOW):
        s6.shapes.add_picture(IMG_FLOW, Inches(1.35), Inches(2.05), width=Inches(4.3), height=Inches(4.5))

    cap_box = s6.shapes.add_textbox(Inches(0.6), Inches(6.62), Inches(5.8), Inches(0.35))
    tf_c = cap_box.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "Figure 6.2: Final Editorial Workflow & Deterministic State Machine"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_MUTED
    p.alignment = PP_ALIGN.CENTER

    c_right = add_card(s6, Inches(6.6), Inches(1.95), Inches(6.133), Inches(5.2))

    t_pipe = s6.shapes.add_textbox(Inches(6.85), Inches(2.10), Inches(5.6), Inches(0.35))
    tf_pipe = t_pipe.text_frame
    pp = tf_pipe.paragraphs[0]
    pp.text = "EXECUTION PIPELINE (4 Key Steps)"
    pp.font.size = Pt(15)
    pp.font.bold = True
    pp.font.color.rgb = NAVY

    pipeline_steps = [
        ("Step 1: Reporter Draft Submission",
         "Field journalist composes news using Formik/Yup editor, tags categories, and uploads cover media to Firebase CDN (status: 'pending').",
         BLUE),
        ("Step 2: Editorial Review & Decision",
         "Chief Editor inspects draft from pending queue, attaches revision notes, or approves article via secure JWT admin endpoint.",
         NAVY),
        ("Step 3: Database Commit & Indexing",
         "PostgreSQL 16 atomically updates status to 'published' and updates B-Tree indexes in under 28ms without table locks.",
         GREEN),
        ("Step 4: Real-Time Sync & Web Push",
         "Socket.IO synchronizes public news feeds across open tabs instantly, while VAPID Web Push alerts subscribers in under 1 second.",
         GOLD)
    ]

    for idx, (st_title, st_desc, st_color) in enumerate(pipeline_steps):
        s_y = Inches(2.55) + idx * Inches(1.05)
        s_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), s_y, Inches(5.6), Inches(0.95))
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = LIGHT_BG
        s_box.line.color.rgb = CARD_BORDER
        s_box.line.width = Pt(1)

        tb = s6.shapes.add_textbox(Inches(7.0), s_y + Inches(0.06), Inches(5.35), Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = st_title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = st_color

        p2 = tf.add_paragraph()
        p2.text = st_desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(2)

    # =========================================================================
    # SLIDE 7: SYSTEM ANALYSIS & DESIGN
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_top_banner(s7)
    add_slide_header_block(s7, 7, "SYSTEM ANALYSIS & DESIGN", "Relational Schema Modeling & Cloud Infrastructure Integration")

    c_left = add_card(s7, Inches(0.6), Inches(1.95), Inches(6.8), Inches(5.2))
    if os.path.exists(IMG_ARCH):
        s7.shapes.add_picture(IMG_ARCH, Inches(0.8), Inches(2.1), width=Inches(6.4), height=Inches(4.35))

    cap_box = s7.shapes.add_textbox(Inches(0.6), Inches(6.60), Inches(6.8), Inches(0.35))
    tf_c = cap_box.text_frame
    tf_c.word_wrap = True
    p = tf_c.paragraphs[0]
    p.text = "Figure 6.1: Final 3-Tier Cloud System Architecture"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = TEXT_MUTED
    p.alignment = PP_ALIGN.CENTER

    c_right = add_card(s7, Inches(7.6), Inches(1.95), Inches(5.133), Inches(5.2))

    design_points = [
        ("9 Relational Tables (ACID)", [
            "• Normalized schemas for users, news, videos, ads, leaves & push.",
            "• Strict foreign keys (ON DELETE CASCADE) prevent data drift."
        ], BLUE),
        ("Secondary B-Tree Indexing", [
            "• Applied to slug, category, status, and published_at DESC.",
            "• Reduced database query latency by >90% (from 320ms to 28ms)."
        ], GREEN),
        ("Atomic Operations & Caching", [
            "• Atomic SQL view counter (views = views + 1) prevents race conditions.",
            "• TanStack React Query placeholder caching eliminates UI flickering."
        ], NAVY)
    ]

    for idx, (title, points, color) in enumerate(design_points):
        top_y = Inches(2.15) + idx * Inches(1.6)
        card_d = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), top_y, Inches(4.733), Inches(1.48))
        card_d.fill.solid()
        card_d.fill.fore_color.rgb = LIGHT_BG
        card_d.line.color.rgb = CARD_BORDER
        card_d.line.width = Pt(1)

        tb = s7.shapes.add_textbox(Inches(7.95), top_y + Inches(0.1), Inches(4.45), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = color

        for pt in points:
            pb = tf.add_paragraph()
            pb.text = pt
            pb.font.size = Pt(12)
            pb.font.color.rgb = TEXT_DARK
            pb.space_before = Pt(4)

    # =========================================================================
    # SLIDE 8: SYSTEM IN ACTION: CONSUMER PORTAL
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_top_banner(s8)
    add_slide_header_block(s8, 8, "SYSTEM IN ACTION: CONSUMER PORTAL", "Public News Web Portal & High-Speed Short-Video Reels")

    c_left = add_card(s8, Inches(0.6), Inches(1.95), Inches(5.9), Inches(5.2))
    tb_l = s8.shapes.add_textbox(Inches(0.8), Inches(2.05), Inches(5.5), Inches(0.35))
    tf_l = tb_l.text_frame
    p = tf_l.paragraphs[0]
    p.text = "Public News Portal & Voice Search"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY

    if os.path.exists(IMG_HERO):
        s8.shapes.add_picture(IMG_HERO, Inches(0.8), Inches(2.45), width=Inches(5.5), height=Inches(2.6))

    tb_b1 = s8.shapes.add_textbox(Inches(0.8), Inches(5.20), Inches(5.5), Inches(1.8))
    tf_b1 = tb_b1.text_frame
    tf_b1.word_wrap = True
    p1 = tf_b1.paragraphs[0]
    p1.text = "• Multi-Lingual Feeds: Native Gujarati (gu) & English (en) with live tickers."
    p1.font.size = Pt(12)
    p1.font.color.rgb = TEXT_DARK

    p2 = tf_b1.add_paragraph()
    p2.text = "• Voice Search: Browser HTML5 Web Speech API converts speech to query text."
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_DARK
    p2.space_before = Pt(4)

    p3 = tf_b1.add_paragraph()
    p3.text = "• Zero UI Flicker: TanStack React Query serves instant cached articles."
    p3.font.size = Pt(12)
    p3.font.color.rgb = TEXT_DARK
    p3.space_before = Pt(4)

    c_right = add_card(s8, Inches(6.8), Inches(1.95), Inches(5.9), Inches(5.2))
    tb_r = s8.shapes.add_textbox(Inches(7.0), Inches(2.05), Inches(5.5), Inches(0.35))
    tf_r = tb_r.text_frame
    p_r = tf_r.paragraphs[0]
    p_r.text = "Short-Video Reels Streamer"
    p_r.font.size = Pt(14)
    p_r.font.bold = True
    p_r.font.color.rgb = NAVY

    if os.path.exists(IMG_VIDEO):
        # 679x658 -> square ratio, place nicely centered
        s8.shapes.add_picture(IMG_VIDEO, Inches(8.3), Inches(2.45), height=Inches(2.6))

    tb_b2 = s8.shapes.add_textbox(Inches(7.0), Inches(5.20), Inches(5.5), Inches(1.8))
    tf_b2 = tb_b2.text_frame
    tf_b2.word_wrap = True
    pr1 = tf_b2.paragraphs[0]
    pr1.text = "• Mobile-First Player: Supports HTML5 MP4 & YouTube Shorts embeds."
    pr1.font.size = Pt(12)
    pr1.font.color.rgb = TEXT_DARK

    pr2 = tf_b2.add_paragraph()
    pr2.text = "• Buffer-Free Playback: <140ms startup latency with swipe gesture controls."
    pr2.font.size = Pt(12)
    pr2.font.color.rgb = TEXT_DARK
    pr2.space_before = Pt(4)

    pr3 = tf_b2.add_paragraph()
    pr3.text = "• Atomic SQL Views: Auto-increments view counts logged directly to database."
    pr3.font.size = Pt(12)
    pr3.font.color.rgb = TEXT_DARK
    pr3.space_before = Pt(4)

    # =========================================================================
    # SLIDE 9: SYSTEM IN ACTION: NEWSROOM GOVERNANCE
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_top_banner(s9)
    add_slide_header_block(s9, 9, "SYSTEM IN ACTION: NEWSROOM GOVERNANCE", "Chief Editor Control Desk & Accredited Reporter Workspace")

    c_left = add_card(s9, Inches(0.6), Inches(1.95), Inches(5.9), Inches(5.2))
    tb_l = s9.shapes.add_textbox(Inches(0.8), Inches(2.05), Inches(5.5), Inches(0.35))
    tf_l = tb_l.text_frame
    p = tf_l.paragraphs[0]
    p.text = "Chief Editor & Admin Control Console"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY

    if os.path.exists(IMG_ADMIN):
        s9.shapes.add_picture(IMG_ADMIN, Inches(0.8), Inches(2.45), width=Inches(5.5), height=Inches(2.6))

    tb_b1 = s9.shapes.add_textbox(Inches(0.8), Inches(5.20), Inches(5.5), Inches(1.8))
    tf_b1 = tb_b1.text_frame
    tf_b1.word_wrap = True
    p1 = tf_b1.paragraphs[0]
    p1.text = "• Editorial Review: 1-click approvals, hero story toggles & revision feedback."
    p1.font.size = Pt(12)
    p1.font.color.rgb = TEXT_DARK

    p2 = tf_b1.add_paragraph()
    p2.text = "• Ad Campaign Manager: Approves client banner ads with live click tracking."
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_DARK
    p2.space_before = Pt(4)

    p3 = tf_b1.add_paragraph()
    p3.text = "• Push Broadcast Center: Dispatches instant VAPID alerts to subscribers."
    p3.font.size = Pt(12)
    p3.font.color.rgb = TEXT_DARK
    p3.space_before = Pt(4)

    c_right = add_card(s9, Inches(6.8), Inches(1.95), Inches(5.9), Inches(5.2))
    tb_r = s9.shapes.add_textbox(Inches(7.0), Inches(2.05), Inches(5.5), Inches(0.35))
    tf_r = tb_r.text_frame
    p_r = tf_r.paragraphs[0]
    p_r.text = "Field Reporter Desk & Digital Press Card"
    p_r.font.size = Pt(14)
    p_r.font.bold = True
    p_r.font.color.rgb = NAVY

    if os.path.exists(IMG_PRESS):
        s9.shapes.add_picture(IMG_PRESS, Inches(7.0), Inches(2.45), width=Inches(5.5), height=Inches(2.6))

    tb_b2 = s9.shapes.add_textbox(Inches(7.0), Inches(5.20), Inches(5.5), Inches(1.8))
    tf_b2 = tb_b2.text_frame
    tf_b2.word_wrap = True
    pr1 = tf_b2.paragraphs[0]
    pr1.text = "• Verifiable Press ID: HTML5 Canvas 2x retina pass with scannable QR code."
    pr1.font.size = Pt(12)
    pr1.font.color.rgb = TEXT_DARK

    pr2 = tf_b2.add_paragraph()
    pr2.text = "• Leave Applications: 4-column leave grid with 1-click editor approval status."
    pr2.font.size = Pt(12)
    pr2.font.color.rgb = TEXT_DARK
    pr2.space_before = Pt(4)

    pr3 = tf_b2.add_paragraph()
    pr3.text = "• Monthly Goals: Tracks articles submitted versus total readership reach."
    pr3.font.size = Pt(12)
    pr3.font.color.rgb = TEXT_DARK
    pr3.space_before = Pt(4)

    # =========================================================================
    # SLIDE 10: RESULTS & VERIFICATION
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_top_banner(s10)
    add_slide_header_block(s10, 10, "RESULTS & VERIFICATION", "Tested in Code, Validated in Production Benchmark Audits")

    stat_cards = [
        ("100%", "Automated Tests Passed", "36 / 36 Test Cases Verified", GREEN),
        ("65 ms", "Avg API Response Latency", "28 ms on Indexed Lookups", BLUE),
        ("0.92 s", "Web Push Dispatch Speed", "VAPID Browser Delivery (<1s)", NAVY),
        ("95 / 100", "Google Lighthouse UX Score", "Zero Mobile Horizontal Overflow", GOLD)
    ]

    for idx, (stat, title, sub, col) in enumerate(stat_cards):
        c_x = Inches(0.6) + idx * Inches(3.08)
        sc = add_card(s10, c_x, Inches(1.95), Inches(2.9), Inches(1.4))

        t_box = s10.shapes.add_textbox(c_x + Inches(0.1), Inches(2.00), Inches(2.7), Inches(1.25))
        tf = t_box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = stat
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = col
        p1.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_DARK
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(2)

        p3 = tf.add_paragraph()
        p3.text = sub
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = TEXT_MUTED
        p3.alignment = PP_ALIGN.CENTER

    c_summary = add_card(s10, Inches(0.6), Inches(3.50), Inches(12.133), Inches(3.65))

    st_box = s10.shapes.add_textbox(Inches(0.85), Inches(3.62), Inches(11.5), Inches(0.35))
    tf_st = st_box.text_frame
    p = tf_st.paragraphs[0]
    p.text = "TESTING & VALIDATION SUMMARY"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY

    qa_points = [
        ("✔ Unit & Module Testing (100% Pass)",
         "Validated Formik/Yup validation, Web Speech dictation hook, Canvas 2x QR generator, and client JPEG compression fallbacks.",
         BLUE),
        ("✔ Concurrency Load Testing",
         "Sustained 500 concurrent read requests maintaining 65ms avg latency with 0% connection drops via PostgreSQL connection pooling.",
         GREEN),
        ("✔ Static TypeScript Audit",
         "0 compilation errors across entire monorepo verified via npx tsc --noEmit for guaranteed runtime type safety.",
         NAVY),
        ("✔ Security & Role Isolation",
         "Strict role route guards (Admin vs Reporter), SQL injection immunity via parameterized queries, CORS headers, and VAPID key signing.",
         GOLD)
    ]

    for idx, (title, desc, col) in enumerate(qa_points):
        row_y = Inches(4.05) + idx * Inches(0.72)
        r_box = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), row_y, Inches(11.633), Inches(0.64))
        r_box.fill.solid()
        r_box.fill.fore_color.rgb = LIGHT_BG
        r_box.line.color.rgb = CARD_BORDER
        r_box.line.width = Pt(1)

        tb = s10.shapes.add_textbox(Inches(1.0), row_y + Inches(0.04), Inches(11.3), Inches(0.55))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title + ": "
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = col

        run = p1.add_run()
        run.text = desc
        run.font.size = Pt(11.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 11: LAUNCHED & OPERATIONAL
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_top_banner(s11)
    add_slide_header_block(s11, 11, "LAUNCHED & OPERATIONAL", "Live on Vercel Cloud CDN & Neon Managed PostgreSQL")

    c_left = add_card(s11, Inches(0.6), Inches(1.95), Inches(7.5), Inches(5.2))

    t_box = s11.shapes.add_textbox(Inches(0.85), Inches(2.15), Inches(7.0), Inches(0.40))
    tf = t_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Enterprise Digital Media Publishing & Streaming Platform"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY

    ops = [
        ("• Cloud Platform", "Vercel Serverless Edge CDN with automated CI/CD from GitHub monorepo."),
        ("• Database Backend", "Managed Neon Cloud PostgreSQL 16 with SSL pooling & B-Tree indexes."),
        ("• Real-Time Engine", "Socket.IO v4.8 WebSockets on Port 3000 & native W3C Web Push alerts."),
        ("• Evaluation Ready", "Active accounts configured for Super Admin, Chief Editor, & Reporter testing."),
        ("• Monorepo Launcher", "Single command (node start.js) concurrently launches all 3 tiers.")
    ]

    for idx, (label, val) in enumerate(ops):
        b_y = Inches(2.70) + idx * Inches(0.72)
        b_shape = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), b_y, Inches(7.0), Inches(0.64))
        b_shape.fill.solid()
        b_shape.fill.fore_color.rgb = LIGHT_BG
        b_shape.line.color.rgb = CARD_BORDER
        b_shape.line.width = Pt(1)

        tb = s11.shapes.add_textbox(Inches(1.0), b_y + Inches(0.04), Inches(6.7), Inches(0.55))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = label + ": "
        p1.font.size = Pt(12)
        p1.font.bold = True
        p1.font.color.rgb = BLUE

        run = p1.add_run()
        run.text = val
        run.font.size = Pt(11.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_DARK

    repo_box = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(6.35), Inches(7.0), Inches(0.65))
    repo_box.fill.solid()
    repo_box.fill.fore_color.rgb = LIGHT_BLUE
    repo_box.line.color.rgb = BLUE
    repo_box.line.width = Pt(1)
    tf_rep = repo_box.text_frame
    tf_rep.word_wrap = True
    p_rep = tf_rep.paragraphs[0]
    p_rep.text = "🌐 Repository: https://github.com/jasanim99-spec/top-newsss-project.git"
    p_rep.font.size = Pt(11)
    p_rep.font.bold = True
    p_rep.font.color.rgb = NAVY
    p_rep.alignment = PP_ALIGN.CENTER

    c_right = add_card(s11, Inches(8.35), Inches(1.95), Inches(4.383), Inches(5.2))
    tb_qr = s11.shapes.add_textbox(Inches(8.55), Inches(2.15), Inches(3.983), Inches(0.40))
    tf_qr = tb_qr.text_frame
    p_q = tf_qr.paragraphs[0]
    p_q.text = "SCAN FOR REPOSITORY"
    p_q.font.size = Pt(14)
    p_q.font.bold = True
    p_q.font.color.rgb = BLUE
    p_q.alignment = PP_ALIGN.CENTER

    if os.path.exists(IMG_QR):
        s11.shapes.add_picture(IMG_QR, Inches(9.10), Inches(2.70), width=Inches(2.9), height=Inches(2.9))

    tb_sub = s11.shapes.add_textbox(Inches(8.55), Inches(5.80), Inches(3.983), Inches(1.10))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True
    p_s = tf_sub.paragraphs[0]
    p_s.text = "Scan with smartphone camera to access source code, documentation & live deployments."
    p_s.font.size = Pt(11.5)
    p_s.font.color.rgb = TEXT_MUTED
    p_s.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 12: ENGINEERING CHALLENGES & SOLUTIONS
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_top_banner(s12)
    add_slide_header_block(s12, 12, "ENGINEERING CHALLENGES & SOLUTIONS", "Key Technical Obstacles Overcome During Development")

    c_main = add_card(s12, Inches(0.6), Inches(1.95), Inches(12.133), Inches(5.2))

    challenges = [
        ("1. Dual-DB Fallback & UI Flickering",
         "Problem: Stale Firestore fallback caused admin dashboard flickering.",
         "Solution: Designated PostgreSQL 16 as single source of truth + React Query.",
         "Result: 100% stable, zero-flicker UI during background refetches.",
         NAVY),

        ("2. Field Media Upload Timeouts",
         "Problem: Large camera photos (>5MB) timed out over weak 3G/4G connections.",
         "Solution: Developed client-side JPEG compression in storageService.ts.",
         "Result: Reduced payload by ~80% & eliminated mobile upload timeouts.",
         NAVY),

        ("3. Cross-Tab View Counter Desync",
         "Problem: Rapid concurrent clicks produced inconsistent view counts.",
         "Solution: Implemented atomic SQL updates (views = views + 1) + Socket.IO.",
         "Result: Accurate real-time view counts synchronized across open client tabs.",
         NAVY),

        ("4. Multi-Service Startup Conflicts",
         "Problem: Running 3 separate servers manually caused port collisions and crashes.",
         "Solution: Built custom start.js orchestrator using child_process.spawn.",
         "Result: Single-command startup with centralized lifecycle management.",
         NAVY)
    ]

    for idx, (title, prob, sol, res, col) in enumerate(challenges):
        row = idx // 2
        col_idx = idx % 2
        c_x = Inches(0.85) + col_idx * Inches(6.0)
        c_y = Inches(2.15) + row * Inches(2.45)
        card_w = Inches(5.7)
        card_h = Inches(2.28)

        ch_card = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, c_y, card_w, card_h)
        ch_card.fill.solid()
        ch_card.fill.fore_color.rgb = LIGHT_BG
        ch_card.line.color.rgb = CARD_BORDER
        ch_card.line.width = Pt(1)

        t_box = s12.shapes.add_textbox(c_x + Inches(0.2), c_y + Inches(0.12), card_w - Inches(0.4), card_h - Inches(0.24))
        tf = t_box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = NAVY

        p_prob = tf.add_paragraph()
        p_prob.text = "• " + prob
        p_prob.font.size = Pt(11)
        p_prob.font.color.rgb = RED
        p_prob.space_before = Pt(3)

        p_sol = tf.add_paragraph()
        p_sol.text = "• " + sol
        p_sol.font.size = Pt(11)
        p_sol.font.color.rgb = TEXT_DARK
        p_sol.space_before = Pt(2)

        p_res = tf.add_paragraph()
        p_res.text = "✔ " + res
        p_res.font.size = Pt(11.5)
        p_res.font.bold = True
        p_res.font.color.rgb = GREEN
        p_res.space_before = Pt(3)

    # =========================================================================
    # SLIDE 13: FUTURE SCOPE & ROADMAP
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13)
    add_top_banner(s13)
    add_slide_header_block(s13, 13, "FUTURE SCOPE & ROADMAP", "Where TOP NEWS Goes Next: Scalability & Expansion")

    c_main = add_card(s13, Inches(0.6), Inches(1.95), Inches(12.133), Inches(5.2))

    futures = [
        ("#1", "Native Mobile Apps (React Native)",
         "Compile cross-platform iOS and Android binaries reusing Express REST APIs, with native APNs/FCM push channels.",
         BLUE),
        ("#2", "Redis In-Memory Caching",
         "Introduce Redis cache with event-driven invalidation to reduce trending feed query latency to sub-5ms.",
         GREEN),
        ("#3", "Neural Text-to-Speech (TTS)",
         "Automatically generate regional audio news podcasts in Gujarati and Hindi for hands-free listening.",
         GOLD),
        ("#4", "Offline-First PWA Mode",
         "Deploy comprehensive offline reading with IndexedDB persistence and Google Workbox service workers.",
         NAVY),
        ("#5", "Automated AI Newsroom Tools",
         "Integrate LLMs for automated article summarization, duplicate detection, and content moderation.",
         BLUE),
        ("#6", "Hyper-Local Citizen Journalism",
         "Launch an eyewitness news tip portal with geo-location tags and programmatic ad exchanges.",
         GREEN)
    ]

    for idx, (num, title, desc, color) in enumerate(futures):
        row = idx // 3
        col = idx % 3
        c_x = Inches(0.85) + col * Inches(3.95)
        c_y = Inches(2.15) + row * Inches(2.45)
        card_w = Inches(3.8)
        card_h = Inches(2.28)

        card_o = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, c_y, card_w, card_h)
        card_o.fill.solid()
        card_o.fill.fore_color.rgb = LIGHT_BG
        card_o.line.color.rgb = CARD_BORDER
        card_o.line.width = Pt(1)

        pill = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x + Inches(0.15), c_y + Inches(0.15), Inches(0.7), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = color
        pill.line.fill.background()
        tf_p = pill.text_frame
        p_p = tf_p.paragraphs[0]
        p_p.text = num
        p_p.font.size = Pt(11)
        p_p.font.bold = True
        p_p.font.color.rgb = CARD_BG
        p_p.alignment = PP_ALIGN.CENTER

        t_head = s13.shapes.add_textbox(c_x + Inches(0.95), c_y + Inches(0.12), card_w - Inches(1.05), Inches(0.38))
        tf_h = t_head.text_frame
        tf_h.word_wrap = True
        ph = tf_h.paragraphs[0]
        ph.text = title
        ph.font.size = Pt(13)
        ph.font.bold = True
        ph.font.color.rgb = NAVY

        t_body = s13.shapes.add_textbox(c_x + Inches(0.15), c_y + Inches(0.60), card_w - Inches(0.3), card_h - Inches(0.70))
        tf_b = t_body.text_frame
        tf_b.word_wrap = True
        pb = tf_b.paragraphs[0]
        pb.text = desc
        pb.font.size = Pt(12)
        pb.font.color.rgb = TEXT_DARK
        pb.space_before = Pt(2)

    # =========================================================================
    # SLIDE 14: CONCLUSION & SUMMARY
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14)
    add_top_banner(s14)
    add_slide_header_block(s14, 14, "CONCLUSION & SUMMARY", "Modernizing Digital Media with High-Performance Full-Stack Engineering")

    c_left = add_card(s14, Inches(0.6), Inches(1.95), Inches(7.5), Inches(5.2))

    pillars = [
        ("High-Speed Decoupled Architecture",
         "Replaced legacy monolithic CMS bottlenecks with React 18, Node.js Express, and Neon PostgreSQL 16, delivering sub-65ms API speeds and zero UI flicker.",
         BLUE),
        ("Integrated Newsroom Governance",
         "Unified multi-lingual publishing, short-video reels, editorial review queues, dynamic QR press cards, and commercial ad tracking into one cohesive platform.",
         GREEN),
        ("Production-Ready & Fully Verified",
         "Achieved 100% test pass rate across 36 scenarios, 0 TypeScript compilation errors (tsc --noEmit), and live production deployment on Vercel Edge CDN.",
         NAVY)
    ]

    for idx, (p_title, p_desc, p_color) in enumerate(pillars):
        p_y = Inches(2.15) + idx * Inches(1.22)
        p_box = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), p_y, Inches(7.0), Inches(1.12))
        p_box.fill.solid()
        p_box.fill.fore_color.rgb = LIGHT_BG
        p_box.line.color.rgb = CARD_BORDER
        p_box.line.width = Pt(1)

        tb = s14.shapes.add_textbox(Inches(1.0), p_y + Inches(0.08), Inches(6.7), Inches(0.96))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = p_title
        p1.font.size = Pt(13.5)
        p1.font.bold = True
        p1.font.color.rgb = p_color

        p2 = tf.add_paragraph()
        p2.text = p_desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_DARK
        p2.space_before = Pt(3)

    rep_box = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(5.95), Inches(7.0), Inches(0.95))
    rep_box.fill.solid()
    rep_box.fill.fore_color.rgb = LIGHT_BLUE
    rep_box.line.color.rgb = BLUE
    rep_box.line.width = Pt(1)
    tf_rep = rep_box.text_frame
    tf_rep.word_wrap = True
    p_rep = tf_rep.paragraphs[0]
    p_rep.text = "GitHub Repository: https://github.com/jasanim99-spec/top-newsss-project.git"
    p_rep.font.size = Pt(11)
    p_rep.font.bold = True
    p_rep.font.color.rgb = NAVY
    p_rep2 = tf_rep.add_paragraph()
    p_rep2.text = "Active Ports: Consumer Portal (:8080)  ·  Admin Console (:5173)  ·  Backend API (:3000)"
    p_rep2.font.size = Pt(10.5)
    p_rep2.font.color.rgb = TEXT_MUTED

    c_right = add_card(s14, Inches(8.35), Inches(1.95), Inches(4.383), Inches(5.2))

    ty_box = s14.shapes.add_textbox(Inches(8.55), Inches(2.15), Inches(3.983), Inches(0.9))
    tf_ty = ty_box.text_frame
    tf_ty.word_wrap = True
    p = tf_ty.paragraphs[0]
    p.text = "THANK YOU!"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER

    p2 = tf_ty.add_paragraph()
    p2.text = "Questions & Technical Discussion"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = BLUE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(3)

    c_det = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.55), Inches(3.30), Inches(3.983), Inches(3.60))
    c_det.fill.solid()
    c_det.fill.fore_color.rgb = LIGHT_BG
    c_det.line.color.rgb = CARD_BORDER
    c_det.line.width = Pt(1)
    tf_cd = c_det.text_frame
    tf_cd.word_wrap = True
    tf_cd.margin_left = Inches(0.2)
    tf_cd.margin_top = Inches(0.15)

    p = tf_cd.paragraphs[0]
    p.text = "CANDIDATE & MENTOR DETAILS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BLUE

    p = tf_cd.add_paragraph()
    p.text = "Jasani Meet Pareshbhai"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_before = Pt(4)

    p = tf_cd.add_paragraph()
    p.text = "Enrollment: 23SE02ML077\nBranch: B.Tech CSE (ML & AI), Sem 7\nSchool of Engineering, P P Savani University"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_DARK
    p.space_before = Pt(3)

    p = tf_cd.add_paragraph()
    p.text = "Industry Mentor:\nMs. Priya Trivedi, Octfis Techno LLP"
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(5)

    p = tf_cd.add_paragraph()
    p.text = "Institute Mentor:\nMr. Rakesh Tanwar, Assistant Professor, PPSU"
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(4)

    # Save both files
    prs.save("TOP_NEWS_FINAL_PRESENTATION.pptx")
    prs.save("TOP_NEWS_PROJECT_PRESENTATION.pptx")
    print(f"Presentations saved successfully! Total slides: {len(prs.slides)}")

if __name__ == "__main__":
    build_presentation()
